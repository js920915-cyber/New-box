"""layout.json(HTML 원래 좌표)으로 Adobe Express 편집용 PPTX를 직접 만든다.

- 슬라이드 1080×1350px(4:5), 1px = 9525 EMU
- 텍스트: 줄마다 독립 텍스트 상자, 자동 줄바꿈·자동 크기 끔, 여백 0, 가운데 정렬
- 굵기별 정적 폰트 패밀리(Noto Sans KR W500~W850) 지정, 자간 보존
- 프레임·박스·번호 원·구분선: 편집 가능한 도형 / 픽토그램: 투명 PNG
사용: python3 source/build_pptx.py
"""
import json, os, re, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUILD = os.path.join(ROOT, "adobe-express", "build")
OUT = os.path.join(ROOT, "adobe-express", "pfp-runners-part2_express.pptx")
PX = 9525
ASC = 0.98          # 텍스트 상자 위끝→첫 줄 기준선(em). LibreOffice 렌더 실측 보정
CHROME_ASC = 1.0    # 브라우저 렌더에 쓰인 가변 폰트 hhea ascent
PAD = 40            # 줄 상자 좌우 여유(px). 폰트 대체 시에도 줄바꿈이 생기지 않게
FRAME = dict(x=34, y=34, w=1012, h=1282, r=28, color="D0DCE5", stroke=2)


def E(v):
    return Emu(int(round(v * PX)))


def rgb(css):
    m = re.findall(r"[\d.]+", css)
    return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))


def alpha(css):
    m = re.findall(r"[\d.]+", css)
    return float(m[3]) if len(m) > 3 else 1.0


def set_round(shape, r, w, h):
    shape.adjustments[0] = min(0.5, r / min(w, h))


def no_shadow(shape):
    """테마 스타일(p:style) 제거 + 빈 효과 목록: 앱별 기본 그림자·윤곽 방지"""
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)
    sp = el.spPr
    if sp.find(qn("a:effectLst")) is None:
        sp.append(sp.makeelement(qn("a:effectLst"), {}))


def add_text_line(slide, line, name):
    size = line["maxSize"]
    baseline = line["glyphTop"] + CHROME_ASC * size
    top = baseline - ASC * size
    height = 1.3 * size
    cx = (line["left"] + line["right"]) / 2
    width = line["right"] - line["left"] + PAD * 2
    tb = slide.shapes.add_textbox(E(cx - width / 2), E(top), E(width), E(height))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    bodyPr.set("wrap", "none")
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.0
    p.space_before = p.space_after = Pt(0)
    for rd in line["runs"]:
        r = p.add_run()
        r.text = rd["text"]
        f = r.font
        f.name = f"Noto Sans KR W{rd['weight']}"
        f.size = Pt(rd["size"] * 0.75)
        f.bold = False
        f.italic = False
        f.color.rgb = rgb(rd["color"])
        rPr = r._r.get_or_add_rPr()
        rPr.set("spc", str(int(round(rd["ls"] * 0.75 * 100))))
        rPr.set("lang", "ko-KR")
        rPr.set("altLang", "en-US")
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", f.name)
    return tb


def main():
    layout = json.load(open(os.path.join(BUILD, "layout.json"), encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(1080), E(1350)
    blank = prs.slide_layouts[6]
    for card in layout:
        n = card["id"][1:]
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # 1) 배경 도형(연한 박스·테두리 박스)
        for i, sh in enumerate(card["shapes"]):
            if sh["kind"] == "line":
                ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(sh["x"]), E(sh["y"]), E(sh["x"] + sh["w"]), E(sh["y"]))
                ln.line.color.rgb = rgb(sh["color"])
                ln.line.width = E(sh["stroke"])
                no_shadow(ln)
                ln.name = f"{n}_구분선_{i+1}"
                continue
            kind = MSO_SHAPE.OVAL if sh["kind"] == "ellipse" else MSO_SHAPE.ROUNDED_RECTANGLE
            shp = s.shapes.add_shape(kind, E(sh["x"]), E(sh["y"]), E(sh["w"]), E(sh["h"]))
            if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                set_round(shp, sh["radius"], sh["w"], sh["h"])
            if alpha(sh["fill"]) > 0:
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(sh["fill"])
            else:
                shp.fill.background()
            if sh["border"]:
                shp.line.color.rgb = rgb(sh["border"]["color"]); shp.line.width = E(sh["border"]["w"])
            else:
                shp.line.fill.background()
            no_shadow(shp)
            label = {"badge": "번호원", "cover-art": "그림배경", "panel": "그림배경", "callout": "강조박스", "next": "예고박스", "": "구간막대"}
            shp.name = f"{n}_{label.get((sh['cls'] or '').split()[0] if sh['cls'] else '', '도형')}_{i+1}"
        # 2) 픽토그램 PNG
        for p in card["pics"]:
            pic = s.shapes.add_picture(os.path.join(BUILD, p["id"] + ".png"), E(p["x"]), E(p["y"]), E(p["w"]), E(p["h"]))
            pic.name = f"{n}_도식"
        # 3) 줄별 텍스트
        for i, line in enumerate(card["lines"]):
            role = line["block"].split(".")[-1]
            add_text_line(s, line, f"{n}_{role}_{i+1:02d}")
        # 4) B01 프레임(맨 위, 채우기 없음)
        fr = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(FRAME["x"]), E(FRAME["y"]), E(FRAME["w"]), E(FRAME["h"]))
        set_round(fr, FRAME["r"], FRAME["w"], FRAME["h"])
        fr.fill.background()
        fr.line.color.rgb = RGBColor.from_string(FRAME["color"])
        fr.line.width = E(FRAME["stroke"])
        no_shadow(fr)
        fr.name = f"{n}_B01프레임"
    prs.core_properties.title = "PFP 2탄 — 각도별 OKC·CKC와 싱글 레그 데드리프트"
    prs.save(OUT)
    print(OUT, os.path.getsize(OUT) // 1024, "KB")


if __name__ == "__main__":
    main()
